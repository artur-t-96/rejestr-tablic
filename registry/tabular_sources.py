"""Odczyt źródła CSV/XLSX bez wykonywania formuł i bez utraty pochodzenia."""

import base64
import csv
import hashlib
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import PurePosixPath
from zipfile import ZipFile

from defusedxml.ElementTree import fromstring
from django.core.exceptions import ValidationError
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

MAX_BYTES = 2_000_000
MAX_UNPACKED_BYTES = 16_000_000
MAX_PART_BYTES = 8_000_000


def _xml(archive, name):
    return fromstring(archive.read(name), forbid_dtd=True, forbid_entities=True, forbid_external=True)


def _xlsx_text(data, sheet_name, fields, date_fields, max_rows):
    try:
        with ZipFile(BytesIO(data)) as archive:
            parts = archive.infolist()
            names = [part.filename for part in parts]
            if (
                len(parts) > 256
                or len(names) != len(set(names))
                or sum(p.file_size for p in parts) > MAX_UNPACKED_BYTES
            ):
                raise ValidationError("XLSX przekracza limit części lub rozpakowanych danych.")
            for part in parts:
                path = PurePosixPath(part.filename)
                if (
                    part.file_size > MAX_PART_BYTES
                    or part.flag_bits & 1
                    or path.is_absolute()
                    or ".." in path.parts
                ):
                    raise ValidationError("XLSX zawiera zbyt dużą lub niedozwoloną część.")
                if any(
                    x in part.filename.lower()
                    for x in ["vbaproject", "activex", "embeddings/", "externallinks/"]
                ):
                    raise ValidationError(
                        "Import nie obsługuje makr, osadzonych obiektów ani połączeń zewnętrznych."
                    )
            types = _xml(archive, "[Content_Types].xml")
            if not any(
                e.get("ContentType")
                == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
                for e in types
            ):
                raise ValidationError("Wybierz zwykły plik XLSX bez makr.")
            workbook = _xml(archive, "xl/workbook.xml")
            relationships = _xml(archive, "xl/_rels/workbook.xml.rels")
            sheets = list(workbook.find("{*}sheets") or [])
            available = [s.get("name") for s in sheets]
            if not sheet_name:
                if len(sheets) != 1:
                    raise ValidationError(
                        "Plik ma kilka arkuszy. Podaj dokładną nazwę wybranego arkusza: "
                        + ", ".join(available[:10])
                    )
                sheet_name = available[0]
            selected = next((s for s in sheets if s.get("name") == sheet_name), None)
            if selected is None or selected.get("state", "visible") != "visible":
                raise ValidationError("Wybrany arkusz nie istnieje lub jest ukryty.")
            rel_id = selected.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            rel = next((r for r in relationships if r.get("Id") == rel_id), None)
            if rel is None:
                raise ValidationError("Brak części wybranego arkusza.")
            target = rel.get("Target", "")
            sheet_path = target.lstrip("/") if target.startswith("/") else "xl/" + target
            if ".." in PurePosixPath(sheet_path).parts:
                raise ValidationError("Niepoprawny adres arkusza.")
            numeric_values = {}
            for name in names:
                if name.endswith((".xml", ".rels")):
                    tree = _xml(archive, name)
                    if name.endswith(".rels") and any(e.get("TargetMode") == "External" for e in tree):
                        raise ValidationError("Usuń odsyłacze i połączenia zewnętrzne przed importem.")
                    if name == sheet_path:
                        # Odrzucamy uszkodzone współrzędne, zanim czytnik mógłby
                        # pominąć wiersz lub zastąpić powtórzoną komórkę.
                        previous_row = 0
                        for row_node in tree.findall("{*}sheetData/{*}row"):
                            row_number = int(row_node.get("r", "0"))
                            if not previous_row < row_number <= max_rows + 1:
                                raise ValidationError("XLSX zawiera niepoprawną kolejność lub numer wiersza.")
                            previous_row = row_number
                            previous_column = 0
                            for cell_node in row_node.findall("{*}c"):
                                column, cell_row = coordinate_from_string(cell_node.get("r", ""))
                                column_number = column_index_from_string(column)
                                if cell_row != row_number or column_number <= previous_column:
                                    raise ValidationError(
                                        "XLSX zawiera powtórzone lub niepoprawne współrzędne komórek."
                                    )
                                previous_column = column_number
                        cells = 0
                        for elem in tree.iter():
                            tag = elem.tag.rsplit("}", 1)[-1]
                            if tag == "mergeCell":
                                raise ValidationError(
                                    "Rozłącz scalone komórki wybranego arkusza przed importem."
                                )
                            if tag in {"row", "col"} and elem.get("hidden") in {"1", "true"}:
                                raise ValidationError(
                                    "Pokaż ukryte wiersze i kolumny wybranego arkusza przed importem."
                                )
                            if tag == "c":
                                cells += 1
                                col, row = coordinate_from_string(elem.get("r", ""))
                                if (
                                    row > max_rows + 1
                                    or column_index_from_string(col) > len(fields)
                                    or cells > (max_rows + 1) * len(fields)
                                ):
                                    raise ValidationError(
                                        f"Arkusz przekracza {max_rows} wierszy danych lub liczbę kolumn wzoru."
                                    )
                                value = elem.find("{*}v")
                                if elem.get("t", "n") == "n" and value is not None:
                                    numeric_values[elem.get("r")] = value.text
                                if elem.find("{*}f") is not None:
                                    raise ValidationError(
                                        f"Komórka {elem.get('r')}: zastąp formułę zweryfikowaną wartością przed importem."
                                    )
            # Wymiary zapisane przez zewnętrzny program nie mogą ucinać danych.
        wb = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            ws = wb[sheet_name]
            ws.reset_dimensions()
            output = StringIO()
            writer = csv.writer(output, delimiter=";")
            header = None
            for index, cells in enumerate(ws.iter_rows(), 1):
                if index > max_rows + 1 or len(cells) > len(fields):
                    raise ValidationError("Arkusz przekracza limit wierszy lub kolumn.")
                if index == 1:
                    header = [c.value for c in cells]
                    if not header or any(not isinstance(v, str) or not v.strip() for v in header):
                        raise ValidationError("Pierwszy wiersz XLSX musi zawierać kompletne nagłówki wzoru.")
                    header = [v.strip() for v in header]
                    writer.writerow(header)
                    continue
                values = []
                if len(cells) > len(header):
                    raise ValidationError(f"Wiersz {index}: dane poza nagłówkami.")
                for col, cell in enumerate(cells):
                    value = cell.value
                    if value is None:
                        value = ""
                    elif cell.data_type in {"f", "e", "b"}:
                        raise ValidationError(
                            f"Komórka {cell.coordinate}: nie importujemy formuł, błędów ani wartości logicznych."
                        )
                    elif isinstance(value, (date, datetime)) and header[col] in date_fields:
                        raw = numeric_values.get(cell.coordinate)
                        if raw and wb.epoch.year == 1899 and Decimal(raw) == 60:
                            raise ValidationError(
                                f"Komórka {cell.coordinate}: nieistniejąca data Excela 1900-02-29. Uzgodnij źródło."
                            )
                        if isinstance(value, datetime) and value.time().isoformat() != "00:00:00":
                            raise ValidationError(
                                f"Komórka {cell.coordinate}: data zawiera godzinę; uzgodnij pełną datę źródłową."
                            )
                        value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
                    elif not isinstance(value, str):
                        raise ValidationError(
                            f"Komórka {cell.coordinate}: zapisz identyfikatory i tekst jako tekst, a daty jako daty lub RRRR-MM-DD."
                        )
                    if len(value) > 32767:
                        raise ValidationError(f"Komórka {cell.coordinate}: zbyt długa wartość.")
                    values.append(value)
                values += [""] * (len(header) - len(values))
                if not any(values):
                    raise ValidationError(
                        f"Wiersz {index}: pusty wiersz wewnątrz wykazu; usuń go przed importem."
                    )
                writer.writerow(values)
            return output.getvalue(), sheet_name
        finally:
            wb.close()
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError(
            "Nie można odczytać poprawnego XLSX. Wybierz niezaszyfrowany plik zgodny ze wzorem."
        ) from exc


def read_source(source, filename, sheet_name, *, fields, date_fields, max_rows):
    data = source.encode("utf-8") if isinstance(source, str) else source
    if not isinstance(data, bytes) or len(data) > MAX_BYTES:
        raise ValidationError("Wybierz CSV UTF-8 lub XLSX do 2 MB.")
    extension = PurePosixPath(filename).suffix.lower()
    if extension == ".xlsx":
        text, selected = _xlsx_text(data, sheet_name, fields, date_fields, max_rows)
        kind = "XLSX"
    elif extension in {"", ".csv"}:
        if sheet_name:
            raise ValidationError("Nazwę arkusza podaje się tylko dla XLSX.")
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise ValidationError("Plik CSV musi być zapisany w UTF-8.") from exc
        selected, kind = "", "CSV"
    else:
        raise ValidationError(
            "Wybierz CSV lub XLSX. Starszy XLS i pliki z makrami wymagają przygotowania do tego formatu."
        )
    return {
        "text": text,
        "sha256": hashlib.sha256(data).hexdigest(),
        "source_format": kind,
        "source_sheet": selected,
        "normalized_csv_sha256": hashlib.sha256(text.encode()).hexdigest(),
    }


def pack_source(data, metadata):
    if metadata["source_format"] == "CSV":
        return {"text": data.decode("utf-8"), "source_format": "CSV", "source_sheet": ""}
    return {
        "source_base64": base64.b64encode(data).decode("ascii"),
        "source_format": "XLSX",
        "source_sheet": metadata["source_sheet"],
    }


def unpack_source(saved):
    if saved.get("source_format", "CSV") == "CSV":
        return saved["text"]
    try:
        return base64.b64decode(saved["source_base64"], validate=True)
    except (ValueError, KeyError) as exc:
        raise ValidationError("Źródło podglądu jest niepoprawne. Wczytaj plik ponownie.") from exc
